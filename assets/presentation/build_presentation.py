"""Build the 18-slide workshop from the agreed story and AGIFORS format reference.

Run from anywhere with python-pptx, Pillow, and qrcode installed. The local
.cache/presentation-python directory is used when it exists. All diagrams,
tables, labels, and formulas are native editable PowerPoint objects.
"""
from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path
import json
import math
import re
import sys

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT / '.cache/presentation-python'))

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches, Pt

OUT = ROOT / 'SCTC-2026-Airlines-Optimization-Workshop.pptx'
NAVY = '0F2343'
DARK = '07142B'
ORANGE = 'F97316'
PALE = 'F7F8FA'
WHITE = 'FFFFFF'
GRAY = '64748B'
LINE = 'DDE4ED'
LIGHT = 'E6EBF2'
SOFT_ORANGE = 'FFF0E5'
GREEN = 'A5D6A7'
GREEN_DARK = '2E7D32'
YELLOW = 'FFF59D'
YELLOW_DARK = '8A6500'
BLUE = '90CAF9'
BLUE_DARK = '1565C0'
W, H = 13.333333, 7.5
ASVG = 'http://schemas.microsoft.com/office/drawing/2016/SVG/main'
REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'

manifest = json.loads((HERE / 'sources.json').read_text(encoding='utf-8'))
template = Path(manifest['format_template'])
if template.exists():
    prs = Presentation(template)
    for item in list(prs.slides._sldIdLst):
        prs.part.drop_rel(item.rId)
        prs.slides._sldIdLst.remove(item)
else:
    prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
prs.core_properties.title = 'Airlines Optimization Workshop: Industry best practices for a decision-support system'
prs.core_properties.subject = 'SCTC 2026: from business requests to trustworthy decision-support software'
prs.core_properties.author = 'Francisco Zenteno Smith'
prs.core_properties.keywords = 'SEFOP, SCTC 2026, operations research, software engineering, validation, verification'
prs.core_properties.comments = '18-slide workshop; visual decisions and presenter material: assets/slide-story.md'

story = (ROOT / 'assets/slide-story.md').read_text(encoding='utf-8')
sections = {}
for match in re.finditer(r'(?ms)^### (\d+)\. (.+?)\n(.*?)(?=^### |^## |\Z)', story):
    sections[int(match.group(1))] = (match.group(2), match.group(3).strip())
planner_script = story.split("## Planner's script (presenter only)")[1].split('## Demonstration boundaries')[0].strip()

timing = {
    1:'0–3', 2:'3–7', 3:'7–10', 4:'10–12', 5:'12–14', 6:'14–15',
    7:'15–29', 8:'29–35', 9:'35–38', 10:'38–42', 11:'42–46', 12:'46–51',
    13:'51–59', 14:'59–65', 15:'65–71', 16:'71–76', 17:'76–80', 18:'80–90',
}

extra_notes = {
    1: 'OPENING: Introduce yourself and the workshop: Francisco Zenteno Smith, Senior Operations Research Scientist, American Airlines; creator of SEFOP. SCTC 2026, Colorado State University Pueblo, October 22, 2026. These views are personal. Explain the three recurring airline decisions. A DSS combines data, business rules and models and repeatedly produces recommendations that a person acts on. Thesis: a model is finished when someone decides with it. No attendee laptop is required. Transition: deciding with it requires being able to trust it.',
    2: 'ACTIVITY: Allow 30–45 seconds of individual thinking, then pair discussion. Ask each pair to write two things they need before acting and one question for the team. Hear contrasting answers and let pairs revise. Do not reveal categories early. Debrief: correct and current data; a model that represents the actual objective and rules; and software that computes that model. Ask: How do you know it is correct, and what did you leave out? Name validation and verification only in the transition to slide 3.',
    3: 'POINT TO THE TWO TRANSLATIONS. Validation: business problem to model; right data, objectives and rules; requirements, examples and feedback. Verification: model to code; design and automated checks. Both are needed. The diagrams are simplified to introduce the distinction, not a claim that the surrounding data and reporting code need no tests. Source: SEFOP book, Section 05, Introduction, verification and validation figure.',
    4: 'STORIES: Geoffrey Chang’s laboratory retracted five structural-biology papers after a data-processing error inverted reported structures. The book describes swapped data columns; avoid implying that the visible paper icons are reproductions of the actual articles. Ariane 5 flight 501 was lost shortly after launch when an unhandled numeric conversion overflow in reused software disabled its inertial reference systems. Connect both to the responsibility for software, including scientific software. Sources: Greg Miller, A Scientist’s Nightmare: Software Problem Leads to Five Retractions, Science 314 (5807), 2006; J. L. Lions et al., Ariane 5 Flight 501 Failure: Report by the Inquiry Board, ESA, 1996; SEFOP book Section 03. The photograph is reused from Purdue workshop slide 6. Do not claim these are optimization-model failures; the lesson concerns software responsibility.',
    5: 'MAP: Part I introduces principles and the cargo example. Part II validates the initial request. Part III verifies the first version. Part IV returns to this same cycle for the next request and then discusses skills and resources. Explain that deployment and use matter: a change helps its user once it is in their hands. The loop is a high-level map; feedback can occur throughout it. Its arrows show a journey, not software dependencies.',
    6: 'SETTING ONLY: For every flight, more shipments are waiting than the aircraft can carry. A load planner chooses which ones fly. Introduce no revenue objective, precise constraint, mathematical formula or selected load on this slide. Explain that what the planner wants reaches the team through a manager, which is why the audience must interview the planner next. Image: built-in imagegen, see assets/presentation/image-prompt.txt.',
    7: 'INTERVIEW: Francisco plays the planner and answers only what is asked. Pairs prepare two questions for one minute and ask in turn. Every pair records answers. The six prompts are guidance, not completed requirements. Translate answers into stories: As a <user>, I want <goal>, so that <reason>. Record the agreed stories as the discussion proceeds. Include the output and invalid-data requirements. Customer identity is input data, but customer diversity is not yet an objective. Full planner script follows below.',
    8: 'EXPLAIN BOTH EXAMPLES. In example 1, the most valuable individual shipment fills the aircraft but earns less than Bighorn and Cedar together. In example 2, the 900 kg shipment cannot fly, and the two 300 kg shipments cannot fly together. Atlas alone is the expected choice. Each situation and expected choice is an acceptance criterion. These are invented teaching cases; they are not claims about real customer bookings.',
    9: 'SYMBOLS: i indexes a shipment; x_i=1 means it flies and 0 means it stays; r_i is revenue in USD; w_i is weight in kg; W is the aircraft weight limit in kg. Every expression comes from a discovered business rule. This workshop uses only a hard weight limit and binary whole-shipment choices; do not import the book appendix’s volume constraint, committed freight or pallet-count variables. The next translation is from the paper model to software.',
    10: 'EXPLAIN THE ANALOGY WITH CARE: small physical variations are often handled by engineering tolerances; software can have discontinuous consequences from one character. Software requirements emerge because users do not fully know what they want, users change their minds, and construction reveals unstated constraints. These are tendencies, not a universal claim that civil engineering has no discovery or software has no tolerances. The resulting abilities are learning quickly and changing quickly and safely. Source: SEFOP book Section 03 and Purdue workshop slide 10.',
    11: 'DEFINE MODULE: a piece of code with a boundary and a clear responsibility. Connections represent dependencies; more connections make a change harder to understand and check. The highlighted module is the same in both diagrams. Separate test boxes show that modules with clear boundaries can be checked individually. Modularity does not mean eliminating necessary connections or guaranteeing that every change fits in one module. Tests check promised behavior and make improving the design safer. Source: SEFOP book Section 04, forces-without-with.svg.',
    12: 'READ THE FIGURES: green unit tests check one module in isolation; yellow integration tests check modules with their real collaborators; blue end-to-end tests exercise the whole program. Each highlighted region represents the scope of one example test, not every test the suite contains. The pyramid suggests many fast unit tests, fewer integration tests and a small number of more costly end-to-end tests; it is guidance, not a quota. EVERY LEVEL still follows Arrange, Act, Assert. Explain orally using slide 8 example 1: arrange the 500 kg flight and shipments 500 kg/$3,000, 250 kg/$2,000, 250 kg/$1,800; run the program; assert Bighorn and Cedar, revenue $3,800 and weight 500 kg. A test must compare with an expected answer; merely running the program is not a check. Source: SEFOP book Section 05, unit-test-pyramid.svg and unit-test-levels.svg.',
    13: 'LIVE DEMO: Run python -m cli optimize_cargo SC101 from the prepared environment. Show each stage announcement, then interpret the recommendation: S1–S4, $9,600, 1,000 kg, two customers. Open src/: reading is adapters/json_data_loader.py; preprocessing, optimization and postprocessing are the folders under use_cases/solving/; output is adapters/console_result_writer.py. Explain each module’s single responsibility. Arrows show execution flow rather than dependency direction. Do not write code live. A successful run is not yet evidence of correctness; transition to the tests.',
    14: 'ACTIVITY: Ask pairs to predict whether this prepared check would catch ignoring weight capacity and allowing fractional shipments. Require reasoning before the debrief. The expected answer is Atlas only, $3,000 and 300 kg. A capacity defect can select Atlas plus Cedar at 600 kg even if preprocessing already excludes the 900 kg shipment. A fractional model can add part of Cedar after preprocessing and exceed $3,000; if the oversize shipment remains available, the fractional optimum is $5,000. The exact wrong result depends on where the bug is introduced. The expected-load/revenue assertion detects those behaviors. Example 1 does not distinguish fractional and binary optimal solutions. LIVE DEMO: open tests/integration/test_situations.py and tests/resources/oversize_shipment_left_behind/data.json, then run pytest. These test functions run the real CLI and real solver even though they use the integration marker. Explain that acceptance describes the promised business behavior, whereas unit/integration/end-to-end describes execution scope. Passing tests give evidence for the covered cases, not proof that no bugs exist.',
    15: 'SECOND TURN, NOT A FEATURE DEMO: On SC101, S1–S4 serve Atlas and Bighorn only; Cedar, Dune, Elk and Fremont have nothing. The planner asks to serve more distinct customers within an acceptable revenue sacrifice. A customer counts once even when several shipments fly. Revisit requirements, examples, model, code, tests, local deployment and use. Do not show an implemented diversity model, new command, allowance values or computed alternatives. Aircraft weight remains hard. Modular design and tests make this next turn manageable.',
    16: 'THREE ELEMENTS, IN FRANCISCO’S OWN WORDS: Business skills. Software skills. People skills. Connect each to your own experience and to the audience’s career opportunities. Leave the screen simple and elaborate verbally. No invented personal anecdote, job history, credential or separate career track has been added.',
    17: 'RESOURCES: Read the book: https://github.com/sefop/sefop-training-hub . Run this workshop: https://github.com/sefop/sefop-sctc-2026 . Ask questions: https://github.com/sefop/sefop-training-hub/discussions . The QR opens the training hub. All three resource labels are clickable. The training hub also points to the SEFOP organization’s discussion board at https://github.com/orgs/sefop/discussions . Explain where the lifecycle, design and testing chapters can be found; invite attendees to keep learning and try the program after the workshop.',
    18: 'DISCUSSION, 10 MINUTES: What remains unclear about turning a business request into decision-support software? Where would you apply learning and adaptation in your own project? Leave the QR visible throughout. If questions slow, invite an example from a student project and trace its next request through validation and verification. Close by returning to the cycle: the next business request begins another turn.',
}


def rgb(value):
    return RGBColor.from_string(value)


def box(slide, x, y, w, h, fill=None, stroke=None, radius=False, line_width=1.5, name=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape._element.spPr.append(OxmlElement('a:effectLst'))
    if fill:
        shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
    else:
        shape.fill.background()
    if stroke:
        shape.line.color.rgb = rgb(stroke); shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    if radius:
        shape.adjustments[0] = .12
    if name:
        shape.name = name
    return shape


def ellipse(slide, x, y, w, h, fill=None, stroke=None, line_width=1.5):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    shape._element.spPr.append(OxmlElement('a:effectLst'))
    if fill:
        shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
    else:
        shape.fill.background()
    if stroke:
        shape.line.color.rgb = rgb(stroke); shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    return shape


def line(slide, x1, y1, x2, y2, color=NAVY, width=2, arrow=False):
    sh = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh._element.spPr.append(OxmlElement('a:effectLst'))
    sh.line.color.rgb = rgb(color); sh.line.width = Pt(width)
    if arrow:
        tail = OxmlElement('a:tailEnd'); tail.set('type', 'triangle'); tail.set('w', 'med'); tail.set('len', 'med')
        sh.line._get_or_add_ln().append(tail)
    return sh


def path(slide, points, color=NAVY, width=2, fill=None, closed=False, arrow=False):
    coords = [(round(x*1000), round(y*1000)) for x,y in points]
    builder = slide.shapes.build_freeform(coords[0][0], coords[0][1], scale=Inches(1)/1000)
    builder.add_line_segments(coords[1:], close=closed)
    sh = builder.convert_to_shape()
    sh._element.spPr.append(OxmlElement('a:effectLst'))
    sh.line.color.rgb = rgb(color); sh.line.width = Pt(width)
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if arrow:
        tail = OxmlElement('a:tailEnd'); tail.set('type', 'triangle'); tail.set('w', 'sm'); tail.set('len', 'sm')
        sh.line._get_or_add_ln().append(tail)
    return sh


def text(slide, value, x, y, w, h, size=24, bold=False, color=NAVY, align='left',
         font='Aptos', valign=MSO_ANCHOR.MIDDLE, margin=0, name=None, link=None):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    sh.name = name or ('Text: ' + value.replace('\n', ' ')[:65])
    tf = sh.text_frame; tf.clear(); tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = valign
    for i, value_line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}[align]
        p.space_before = Pt(0); p.space_after = Pt(0); p.line_spacing = 1.06
        run = p.add_run(); run.text = value_line
        run.font.name = font; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = rgb(color)
    if link:
        sh.click_action.hyperlink.address = link
    return sh


def picture(slide, file, x, y, w, h, contain=True, crop=None):
    file = HERE / file
    with Image.open(file) as im:
        iw, ih = im.size
    if contain:
        ratio = min(w / iw, h / ih)
        ww, hh = iw * ratio, ih * ratio
        sh = slide.shapes.add_picture(str(file), Inches(x + (w-ww)/2), Inches(y + (h-hh)/2), width=Inches(ww), height=Inches(hh))
    else:
        sh = slide.shapes.add_picture(str(file), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
        if crop:
            sh.crop_left, sh.crop_top, sh.crop_right, sh.crop_bottom = crop
        else:
            src_ar, target_ar = iw/ih, w/h
            if src_ar > target_ar:
                sh.crop_left = sh.crop_right = (1-target_ar/src_ar)/2
            else:
                sh.crop_top = sh.crop_bottom = (1-src_ar/target_ar)/2
    sh.name = file.stem
    sh._element.nvPicPr.cNvPr.set('descr', file.stem.replace('-', ' '))
    svg = file.with_suffix('.svg')
    if svg.exists():
        partname = prs.part.package.next_partname('/ppt/media/vector%d.svg')
        svg_part = Part(partname, 'image/svg+xml', prs.part.package, blob=svg.read_bytes())
        rid = slide.part.relate_to(svg_part, RT.IMAGE)
        blip = sh._element.blipFill.blip
        extlst = OxmlElement('a:extLst'); ext = OxmlElement('a:ext')
        ext.set('uri', '{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
        from lxml import etree
        svgblip = etree.Element('{'+ASVG+'}svgBlip', nsmap={'asvg':ASVG})
        svgblip.set('{'+REL+'}embed', rid)
        ext.append(svgblip); extlst.append(ext); blip.append(extlst)
    return sh


def icon(slide, kind, x, y, size=.7, color=NAVY, accent=ORANGE):
    """Small native line icons in a consistent geometric vocabulary."""
    def b(xx,yy,ww,hh,fill=None,stroke=color,radius=False,lw=2):
        return box(slide,x+xx*size,y+yy*size,ww*size,hh*size,fill,stroke,radius,lw)
    def e(xx,yy,ww,hh,fill=None,stroke=color,lw=2):
        return ellipse(slide,x+xx*size,y+yy*size,ww*size,hh*size,fill,stroke,lw)
    def l(x1,y1,x2,y2,col=color,lw=2):
        return line(slide,x+x1*size,y+y1*size,x+x2*size,y+y2*size,col,lw)
    if kind in ('person','people'):
        if kind == 'people':
            e(.04,.20,.24,.24, color,None); e(.72,.20,.24,.24,color,None)
            b(0,.51,.33,.35,color,None,True); b(.67,.51,.33,.35,color,None,True)
        e(.34,.05,.32,.32, color,None)
        b(.24,.46,.52,.49,color,None,True)
    elif kind == 'file':
        b(.19,.07,.62,.86,WHITE,color,True)
        l(.31,.34,.69,.34); l(.31,.50,.69,.50); l(.31,.67,.57,.67,accent)
    elif kind == 'model':
        text(slide,'∑',x,y,size,size,46*size,False,color,'center','Cambria Math')
    elif kind == 'code':
        l(.30,.23,.10,.50); l(.10,.50,.30,.77)
        l(.70,.23,.90,.50); l(.90,.50,.70,.77)
        l(.59,.15,.42,.85,accent)
    elif kind == 'test':
        b(.13,.10,.74,.80,WHITE,color,True)
        l(.28,.49,.44,.65,accent,2.5); l(.44,.65,.73,.30,accent,2.5)
    elif kind == 'deploy':
        b(.13,.50,.74,.40,WHITE,color,True)
        l(.50,.10,.50,.64,accent,2.5); l(.28,.32,.50,.10,accent,2.5); l(.72,.32,.50,.10,accent,2.5)
    elif kind in ('chat','feedback'):
        b(.07,.12,.86,.60,WHITE,color,True)
        l(.24,.72,.24,.90); l(.24,.90,.45,.72)
        for xx in (.28,.50,.72): e(xx-.035,.395,.07,.07,accent,None)
    elif kind == 'modules':
        b(.07,.10,.32,.30,WHITE,color,True); b(.61,.10,.32,.30,WHITE,color,True)
        b(.34,.62,.32,.30,WHITE,color,True)
        l(.23,.40,.23,.51); l(.77,.40,.77,.51); l(.23,.51,.77,.51); l(.50,.51,.50,.62,accent)
    elif kind == 'result':
        b(.07,.13,.86,.73,WHITE,color,True)
        l(.12,.34,.88,.34,accent)
        for yy in (.50,.67): l(.19,yy,.80,yy)
        l(.48,.34,.48,.80)
    elif kind == 'book':
        b(.07,.16,.41,.69,WHITE,color,True); b(.52,.16,.41,.69,WHITE,color,True)
        l(.50,.18,.50,.90,accent)
        l(.16,.34,.38,.34); l(.62,.34,.84,.34); l(.16,.50,.38,.50); l(.62,.50,.84,.50)
    elif kind == 'business':
        b(.06,.28,.88,.58,WHITE,color,True); b(.34,.10,.32,.18,WHITE,color,True)
        l(.08,.48,.92,.48); b(.44,.43,.12,.12,accent,None)
    elif kind == 'company':
        b(.18,.11,.64,.82,None,color,False)
        for xx in (.32,.60):
            for yy in (.28,.49): b(xx,yy,.10,.10,color,None)
        b(.43,.73,.14,.20,color,None)


def notes(slide, number):
    title, body = sections[number]
    content = re.sub(r'(?ms)^- Visual:.*?(?=^- [^ ]|\Z)', '', body).strip()
    if number == 4:
        content = content.replace('a sign error in a data-processing script\n  retracted five published papers', 'a data-processing error\n  led to five published papers being retracted')
        content = content.replace('rocket 37 seconds after launch', 'rocket shortly after launch')
    if number == 14:
        content = re.sub(r'(?s)- Answer:.*?(?=- Connecting idea:)', '', content)
    value = f'{number:02d} | {title}\nSuggested elapsed workshop minutes: {timing[number]} (90 minutes total).\n\nPRESENTER GUIDE\n{extra_notes[number]}\n\nSTORY TALKING POINTS\n{content}'
    if number == 7:
        value += '\n\nPLANNER SCRIPT (PRESENTER ONLY)\n' + planner_script
    value += '\n\nWorkshop source: assets/slide-story.md. Format reference: AGIFORS Annual Symposium 2026 deck.\n'
    slide.notes_slide.notes_text_frame.text = value


def new_slide(number, title=None, size=36):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(PALE)
    # Remove inherited layout placeholders; all visible objects below are explicit.
    for placeholder in list(slide.placeholders):
        placeholder._element.getparent().remove(placeholder._element)
    if title:
        title_y, title_h = (.27,1.30) if '\n' in title else (.32,1.15)
        text(slide,title,.55,title_y,12.23,title_h,size,True,NAVY,font='Aptos Display',name='Slide title')
        box(slide,.58,1.60,.58,.055,ORANGE)
    text(slide,'Francisco Zenteno Smith  ·  SEFOP  ·  SCTC 2026',.58,7.10,10.6,.20,11,color=GRAY)
    text(slide,str(number),12.18,7.08,.55,.24,12,color=GRAY,align='right')
    notes(slide,number)
    return slide


def note_band(slide, value, y=6.50, x=.65, w=12.03, size=22):
    box(slide,x,y,w,.50,SOFT_ORANGE,None,True)
    text(slide,value,x+.14,y+.04,w-.28,.42,size,True,NAVY,'center')


def cargo_table(slide, which, x, y, w=5.75, result=True):
    rows = [('Atlas Steel',500,3000),('Bighorn Auto',250,2000),('Cedar Pharma',250,1800)] if which==1 else [('Atlas Steel',300,3000),('Bighorn Auto',900,9000),('Cedar Pharma',300,800)]
    selected = [1,2] if which==1 else [0]
    h_header, h_row = .52,.66
    table = slide.shapes.add_table(4,3,Inches(x),Inches(y),Inches(w),Inches(h_header+3*h_row)).table
    table.columns[0].width=Inches(w*.48); table.columns[1].width=Inches(w*.24); table.columns[2].width=Inches(w*.28)
    table.rows[0].height=Inches(h_header)
    for row in table.rows:
        for cell in row.cells:
            cell.margin_left=Inches(.14); cell.margin_right=Inches(.12)
            cell.margin_top=Inches(.05); cell.margin_bottom=Inches(.05)
            cell.vertical_anchor=MSO_ANCHOR.MIDDLE
    values=[('Customer','Weight','Revenue')]+[(a,f'{b:,} kg',f'${c:,}') for a,b,c in rows]
    for r, items in enumerate(values):
        if r: table.rows[r].height=Inches(h_row)
        for c, value in enumerate(items):
            cell=table.cell(r,c); cell.text=value
            cell.fill.solid(); cell.fill.fore_color.rgb=rgb(NAVY if r==0 else (SOFT_ORANGE if r-1 in selected else WHITE))
            p=cell.text_frame.paragraphs[0]; p.alignment=PP_ALIGN.LEFT if c==0 else PP_ALIGN.RIGHT
            p.space_before=Pt(0); p.space_after=Pt(0)
            for run in p.runs:
                run.font.name='Aptos'; run.font.size=Pt(20 if r else 18); run.font.bold=(r==0 or r-1 in selected)
                run.font.color.rgb=rgb(WHITE if r==0 else NAVY)
            tcpr=cell._tc.get_or_add_tcPr()
            for edge in ['lnL','lnR','lnT','lnB']:
                el=OxmlElement('a:'+edge); el.set('w','6350'); fill=OxmlElement('a:solidFill'); clr=OxmlElement('a:srgbClr'); clr.set('val',LINE); fill.append(clr); el.append(fill); tcpr.append(el)
    for row in selected:
        box(slide,x,y+h_header+row*h_row,w,h_row,None,ORANGE,False,1.6)
        text(slide,'✓',x-.41,y+h_header+row*h_row+.09,.34,.42,21,True,ORANGE,'center','Segoe UI Symbol')
    if result:
        value='Bighorn + Cedar' if which==1 else 'Atlas'
        total='$3,800  ·  500 kg' if which==1 else '$3,000  ·  300 kg'
        text(slide,value,x,y+2.75,w,.42,23,True,align='center')
        text(slide,total,x,y+3.24,w,.47,27,True,align='center')
    return table


def cycle(slide, second=False):
    cx,cy,rx,ry=6.66,4.40,4.43,1.80
    labels=['Request','Requirements','Model','Design & code','Tests','Deploy','Use','Feedback']
    icons=['chat','file','model','code','test','deploy','person','feedback']
    angles=[-90,-45,0,45,90,135,180,225]
    for i, a in enumerate(angles):
        # Clip the arc at the two node boundaries, so its arrowhead stays visible.
        centers=[(cx+rx*math.cos(math.radians(t)),cy+ry*math.sin(math.radians(t))) for t in [a,a+45]]
        pts=[]
        for j in range(201):
            t=math.radians(a+45*j/200)
            p=(cx+rx*math.cos(t),cy+ry*math.sin(t))
            if all(not (abs(p[0]-xx)<1.30 and abs(p[1]-yy)<.43) for xx,yy in centers):
                pts.append(p)
        path(slide,pts,ORANGE if i==7 else GRAY,2.5 if i==7 else 1.7,arrow=True)
    for i,a in enumerate(angles):
        xx=cx+rx*math.cos(math.radians(a)); yy=cy+ry*math.sin(math.radians(a))
        nodew,nodeh=2.50,.76
        box(slide,xx-nodew/2,yy-nodeh/2,nodew,nodeh,WHITE,ORANGE if second and i==0 else LINE,True,2 if second and i==0 else 1.2,name='Lifecycle: '+labels[i])
        icon(slide,icons[i],xx-nodew/2+.10,yy-.24,.44)
        text(slide,labels[i],xx-nodew/2+.61,yy-.30,nodew-.70,.60,18 if i in [1,3] else 20,True)
    text(slide,'Validation',8.70,2.08,2.20,.35,19,True,align='center')
    text(slide,'Verification',8.70,6.35,2.20,.35,19,True,align='center')
    # Small brackets identify the two arcs without changing the workshop's shared colors.
    line(slide,8.50,2.51,10.13,2.51,ORANGE,2)
    line(slide,8.50,2.51,8.50,2.66,ORANGE,2)
    line(slide,10.13,2.51,10.13,2.66,ORANGE,2)
    line(slide,8.50,6.20,10.13,6.20,ORANGE,2)
    line(slide,8.50,6.07,8.50,6.20,ORANGE,2)
    line(slide,10.13,6.07,10.13,6.20,ORANGE,2)
    if second:
        box(slide,4.81,3.10,3.70,.95,SOFT_ORANGE,None,True)
        text(slide,'Can we serve more customers?',4.95,3.18,3.42,.80,21,True,align='center')
        for i in range(6):
            xx=4.86+i*.61
            # Filled and outline customer buildings, not shipment boxes.
            if i<2:
                box(slide,xx,4.28,.43,.62,NAVY,None,True)
                for dx in [.09,.25]:
                    for dy in [.13,.31]:box(slide,xx+dx,4.28+dy,.07,.07,WHITE)
            else:
                box(slide,xx,4.28,.43,.62,None,GRAY,True,1.3)
                for dx in [.09,.25]:
                    for dy in [.13,.31]:box(slide,xx+dx,4.28+dy,.07,.07,GRAY)
        text(slide,'2 of 6 customers served',4.55,5.12,4.18,.43,23,True,align='center')


# 01 — recurring decisions
s=new_slide(1,'A decision-support system supports\nrecurring business decisions')
for cx,file,label,cadence in [(2.62,'fleet.png','Plan the fleet','Annually'),(6.66,'crew.png','Schedule crews','Monthly'),(10.70,'disruption.png','Recover operations','Daily')]:
    picture(s,file,cx-1.0,2.25,2.0,2.0)
    text(s,label,cx-1.91,4.66,3.82,.55,27,True,align='center')
    text(s,cadence,cx-1.7,5.43,3.4,.55,26,True,ORANGE,'center')

# 02 — the trust question
s=new_slide(2,'What would you need before trusting\na recommendation?')
box(s,.96,2.26,4.73,2.83,WHITE,NAVY,True,2.5)
box(s,1.09,2.39,4.47,2.46,WHITE)
text(s,'OPTIMAL',1.24,3.13,4.17,.80,43,True,align='center')
box(s,3.00,5.10,.65,.45,NAVY)
box(s,2.38,5.53,1.88,.13,NAVY,None,True)
text(s,'?',6.20,2.90,.95,1.57,88,True,ORANGE,'center','Aptos Display')
# A deliberately simple, editable person considering the result.
ellipse(s,9.45,2.40,.90,.90,NAVY)
box(s,8.80,3.49,2.20,1.70,NAVY,None,True)
line(s,10.72,4.08,11.28,4.70,NAVY,18)
line(s,11.28,4.70,10.37,3.42,NAVY,18)
ellipse(s,10.17,3.18,.37,.37,ORANGE)
line(s,8.50,5.55,11.60,5.55,GRAY,2)
note_band(s,'Discuss in pairs: two things you need to know; one question for the team.',6.40,size=20)

# 03 — the two translations
s=new_slide(3,'Deploying a tool requires\nvalidation and verification')
for x,label,kind in [(.72,'Business problem','people'),(4.84,'Mathematical model','model'),(8.96,'Code','code')]:
    box(s,x,2.43,3.65,2.05,WHITE,LINE,True)
    icon(s,kind,x+1.28,2.69,1.08)
    text(s,label,x+.10,3.91,3.45,.43,25,True,align='center')
line(s,4.43,3.46,4.78,3.46,ORANGE,2.5,True)
line(s,8.55,3.46,8.90,3.46,ORANGE,2.5,True)
for x,label,question in [(2.58,'Validation','The right problem?'),(6.70,'Verification','Solved right?')]:
    line(s,x,4.75,x+3.65,4.75,ORANGE,2)
    line(s,x,4.56,x,4.75,ORANGE,2);line(s,x+3.65,4.56,x+3.65,4.75,ORANGE,2)
    text(s,label,x-.28,5.07,4.21,.55,29,True,align='center')
    text(s,question,x-.28,5.78,4.21,.44,25,align='center')

# 04 — failures in research and industry
s=new_slide(4,'A decision-support system is software,\nso the rules of software apply')
text(s,'Research',.78,1.95,5.60,.45,27,True,align='center')
text(s,'Industry',6.96,1.95,5.60,.45,27,True,align='center')
for i,(xx,yy) in enumerate([(1.12,2.80),(2.78,2.80),(4.44,2.80),(1.95,4.12),(3.61,4.12)]):
    box(s,xx,yy,1.37,1.55,WHITE,LINE,True)
    for dy,length in [(.28,.83),(.47,.94),(.66,.67)]:line(s,xx+.17,yy+dy,xx+.17+length,yy+dy,LINE,1.8)
    stamp=box(s,xx-.04,yy+.89,1.45,.35,None,ORANGE,False,1.2)
    text(s,'RETRACTED',xx-.03,yy+.90,1.43,.29,12,True,ORANGE,'center')
picture(s,'ariane5.jpg',7.00,2.69,5.50,3.31,contain=False)
text(s,'A data-processing error',.78,6.17,5.60,.45,25,True,align='center')
text(s,'An unhandled overflow',6.96,6.17,5.60,.45,25,True,align='center')

# 05 — the lifecycle
s=new_slide(5,'Working on a DSS is a cycle')
cycle(s)

# 06 — the setting, with no requirements revealed
s=new_slide(6,'Choosing the cargo for each flight')
picture(s,'cargo-planner.png',.57,1.91,12.20,4.91,contain=False)

# 07 — the planner interview
s=new_slide(7,'Interview the planner to find the requirements',size=35)
box(s,1.10,1.95,11.13,.95,WHITE,LINE,True)
text(s,'“I heard our cargo planner wants a tool that picks\nthe best shipments for each flight.”',1.36,2.00,10.61,.86,24,align='center')
picture(s,'cargo-planner.png',5.47,3.25,2.40,2.55,contain=False,crop=(.04,.12,.64,.12))
for x,y,label,kind in [(.95,3.20,'Decisions','chat'),(.95,4.32,'Objectives','business'),(.95,5.44,'Constraints','file'),(9.06,3.20,'Data','result'),(9.06,4.32,'Uncertainty','model'),(9.06,5.44,'Risks','test')]:
    icon(s,kind,x,y+.08,.51)
    text(s,label,x+.74,y,3.10,.71,26,True)
note_band(s,'Pairs  ·  1 minute  ·  2 questions',6.42)

# 08 — acceptance examples
s=new_slide(8,'Examples show what the planner would do')
for x,n in [(.85,1),(6.95,2)]:
    text(s,f'Example {n}',x,1.93,5.50,.42,25,True,align='center')
    text(s,'500 kg limit',x,2.42,5.50,.34,21,True,ORANGE,'center')
    cargo_table(s,n,x,3.02,5.50)

# 09 — business rules translated into mathematics
s=new_slide(9,'The business rules lead to a knapsack model',size=35)
for y,label,formula in [(2.19,'Each shipment flies\nwhole or stays','xᵢ ∈ {0, 1}'),(3.57,'Earn the most revenue','max  ∑ᵢ rᵢ xᵢ'),(4.95,'Stay within the weight limit','∑ᵢ wᵢ xᵢ ≤ W')]:
    text(s,label,.82,y-.04,5.15,1.02,27,True)
    line(s,6.10,y+.46,6.68,y+.46,ORANGE,2.5,True)
    text(s,formula,7.10,y-.03,5.40,.97,38,font='Cambria Math')
text(s,'i: shipment   ·   xᵢ: fly / stay   ·   rᵢ: revenue   ·   wᵢ: weight   ·   W: weight limit',.74,6.43,11.85,.48,19,align='center')

# 10 — different engineering conditions
s=new_slide(10,'Software is a different kind of engineering',size=35)
for x,label,file,a,b in [(.78,'Civil engineering','civil-engineering.jpg','Requirements largely known','Small errors tolerated'),(6.95,'Software engineering','software-engineering.jpg','Requirements emerge','Tiny errors can matter')]:
    text(s,label,x,1.97,5.60,.45,27,True,align='center')
    picture(s,file,x,2.58,5.60,2.66,contain=False)
    text(s,a,x,5.47,5.60,.42,23,True,align='center')
    text(s,b,x,6.00,5.60,.41,23,align='center')
text(s,'Learn quickly  ·  Change quickly and safely',.80,6.62,11.75,.37,23,True,ORANGE,'center')

# 11 — a comparison of the same six modules
s=new_slide(11,'Design and testing make\nquick, safe change possible')
left=[(1.18,2.72),(3.45,2.72),(5.03,3.92),(3.45,5.20),(1.18,5.20),(.78,3.92)]
right=[(7.31,2.85),(9.10,2.85),(10.89,2.85),(7.31,4.69),(9.10,4.69),(10.89,4.69)]
for i in range(6):
    for j in range(i+1,6):
        line(s,left[i][0]+.41,left[i][1]+.30,left[j][0]+.41,left[j][1]+.30,LINE,1.3)
for i,j in [(0,1),(1,2),(0,3),(1,4),(2,5),(3,4),(4,5)]:
    line(s,right[i][0]+.41,right[i][1]+.30,right[j][0]+.41,right[j][1]+.30,GRAY,1.5)
for positions,tests in [(left,False),(right,True)]:
    for i,(x,y) in enumerate(positions):
        box(s,x,y,.82,.60,ORANGE if i==1 else WHITE,ORANGE if i==1 else NAVY,True,1.7)
        text(s,chr(65+i),x,y+.05,.82,.49,22,True,WHITE if i==1 else NAVY,'center')
        if tests:
            line(s,x+.41,y+.60,x+.41,y+1.02,GRAY,1.3)
            box(s,x+.08,y+1.02,.66,.40,WHITE,NAVY,True,1.1)
            text(s,'✓',x+.15,y+1.03,.51,.34,17,True,ORANGE,'center','Segoe UI Symbol')
text(s,'Tests',11.90,5.77,.80,.35,17,color=GRAY)
text(s,'Changes spread',.71,6.38,5.63,.48,27,True,align='center')
text(s,'Changes stay manageable',6.85,6.38,5.70,.48,27,True,align='center')

# 12 — the testing pyramid and execution scopes
s=new_slide(12,'What an automated test is')
left,top,pw,ph=.75,2.45,3.13,3.14
cx=left+pw/2
poly=[([(cx,top),(cx+.47,top+.94),(cx-.47,top+.94)],BLUE,BLUE_DARK),
      ([(cx-.54,top+1.06),(cx+.54,top+1.06),(cx+1.00,top+2.00),(cx-1.00,top+2.00)],YELLOW,YELLOW_DARK),
      ([(cx-1.07,top+2.12),(cx+1.07,top+2.12),(cx+1.55,top+3.06),(cx-1.55,top+3.06)],GREEN,GREEN_DARK)]
for pts,fill,stroke in poly:path(s,pts,stroke,1.4,fill,True)
text(s,'End-to-end',left,1.99,pw,.37,18,True,BLUE_DARK,'center')
line(s,cx,2.36,cx,top,BLUE_DARK,1.3)
text(s,'Integration',left+.27,top+1.42,pw-.54,.36,18,True,YELLOW_DARK,'center')
text(s,'Unit',left+.27,top+2.44,pw-.54,.38,21,True,GREEN_DARK,'center')
for col,(label,colr,stroke) in enumerate([('Unit',GREEN,GREEN_DARK),('Integration',YELLOW,YELLOW_DARK),('End-to-end',BLUE,BLUE_DARK)]):
    x=4.46+col*2.66
    text(s,label,x-.07,2.04,2.47,.43,25,True,stroke,'center')
    if col==1:box(s,x-.13,3.75,2.48,1.78,YELLOW,YELLOW_DARK,True,1.5)
    if col==2:box(s,x-.13,2.81,2.48,2.72,BLUE,BLUE_DARK,True,1.5)
    for row,label2 in enumerate(['Read','Compute','Report']):
        yy=2.99+row*.92
        if row<2:line(s,x+1.10,yy+.57,x+1.10,yy+.87,GRAY,1.7,True)
        box(s,x,yy,2.21,.57,GREEN if col==0 and row==1 else WHITE,stroke if col==0 and row==1 else LINE,True,1.4)
        text(s,label2,x+.05,yy+.06,2.11,.43,22,True,align='center')
for xx,label,caption in [(2.00,'Arrange','Set up'),(5.27,'Act','Run'),(8.53,'Assert','Check')]:
    text(s,label,xx,6.08,2.80,.52,28,True,align='center')
    text(s,caption,xx,6.59,2.80,.38,20,color=GRAY,align='center')
for xx in [4.83,8.10]:text(s,'→',xx,6.12,.40,.44,26,True,ORANGE,'center')

# 13 — match the running program's five stages
s=new_slide(13,'A DSS example design')
labels=['Read data','Preprocess','Solve','Postprocess','Result']
for i,(label,kind) in enumerate(zip(labels,['file','modules','model','modules','result'])):
    x=.60+i*2.55
    box(s,x,2.72,2.12,2.35,WHITE,LINE,True)
    icon(s,kind,x+.58,3.01,.96)
    text(s,label,x+.06,4.37,2.00,.44,22 if len(label)<11 else 21,True,align='center')
    if i<4:line(s,x+2.20,3.88,x+2.47,3.88,ORANGE,2.5,True)
note_band(s,'SC101  ·  Live demonstration',6.14)

# 14 — audience predictions before the live test run
s=new_slide(14,'The planner’s examples become acceptance tests',size=34)
text(s,'500 kg limit',.92,2.15,5.35,.39,23,True,ORANGE,'center')
cargo_table(s,2,.92,2.84,5.35,False)
text(s,'Expected: Atlas',.92,5.54,5.35,.43,25,True,align='center')
text(s,'$3,000  ·  300 kg',.92,6.10,5.35,.48,27,True,align='center')
text(s,'Would this test catch…',6.91,2.16,5.50,.43,26,True)
for y,value in [(3.00,'Ignoring the weight limit?'),(4.57,'Splitting a shipment?')]:
    box(s,6.91,y,5.50,1.15,WHITE,LINE,True)
    text(s,value,7.18,y+.17,4.96,.78,26,True)
text(s,'Predict in pairs  ·  Explain why',6.91,6.25,5.50,.43,21,True,ORANGE,'center')

# 15 — the same lifecycle, next request
s=new_slide(15,'The cycle runs again:\nthe planner asks for client diversity')
cycle(s,True)

# 16 — the three elements the speaker will elaborate
s=new_slide(16,'What this profession needs')
for cx,label,kind in [(2.60,'Business skills','business'),(6.66,'Software skills','modules'),(10.72,'People skills','people')]:
    icon(s,kind,cx-.85,2.56,1.70)
    text(s,label,cx-1.82,5.00,3.64,.61,29,True,align='center')

# 17 — resources with a single primary destination
s=new_slide(17,'SEFOP: where to keep learning')
links=[('Read the book','book','https://github.com/sefop/sefop-training-hub'),('Run the workshop','code','https://github.com/sefop/sefop-sctc-2026'),('Ask questions','chat','https://github.com/sefop/sefop-training-hub/discussions')]
for y,(label,kind,url) in zip([2.55,3.92,5.29],links):
    icon(s,kind,.98,y,.78)
    text(s,label,2.04,y+.03,4.62,.70,29,True,link=url)
picture(s,'training-hub-qr.png',7.51,2.05,4.58,4.58)
text(s,'github.com/sefop/sefop-training-hub',6.88,6.48,5.84,.37,18,align='center',link='https://github.com/sefop/sefop-training-hub')

# 18 — leave this visible during the discussion
s=new_slide(18)
text(s,'Questions?',.83,2.79,6.18,1.18,62,True,font='Aptos Display')
box(s,.89,4.32,.92,.065,ORANGE)
picture(s,'training-hub-qr.png',7.51,1.65,4.58,4.58)
text(s,'github.com/sefop/sefop-training-hub',6.88,6.13,5.84,.37,18,align='center',link='https://github.com/sefop/sefop-training-hub')

assert len(prs.slides)==18
for i,slide in enumerate(prs.slides,1):
    assert slide.has_notes_slide and len(slide.notes_slide.notes_text_frame.text)>200, i

prs.save(OUT)
print(f'Saved {len(prs.slides)} slides with presenter notes: {OUT}')

# Save a compact manifest of visible text for content review.
review=[]
for i,slide in enumerate(prs.slides,1):
    visible=[]
    for shape in slide.shapes:
        if shape.has_text_frame:visible.append(shape.text)
        if shape.has_table:
            visible.extend(' | '.join(cell.text for cell in row.cells) for row in shape.table.rows)
    review.append({'slide':i,'visible_text':visible,'notes_characters':len(slide.notes_slide.notes_text_frame.text)})
(ROOT / '.cache/presentation-review.json').write_text(json.dumps(review,indent=2,ensure_ascii=False),encoding='utf-8')
